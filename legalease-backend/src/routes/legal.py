from flask import Blueprint, jsonify, request, send_file
from flask_cors import cross_origin
import google.generativeai as genai
import os
import uuid
import base64
import tempfile
from fpdf import FPDF
import json
import io

from PyPDF2 import PdfReader
from docx import Document

from src.templates.agreement_templates import get_template_by_type, format_template

legal_bp = Blueprint('legal', __name__)

# Configure Gemini API
genai.configure(api_key=os.environ.get('GEMINI_API_KEY'))
model = genai.GenerativeModel('gemini-1.5-flash')

# In-memory storage for documents and chat history
documents = {}
chat_histories = {}

# ------------------- Helper functions -------------------

def extract_pdf_text(file_bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def extract_docx_text(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    return text

def safe_format(template_str, data):
    """Safely format a string using keys in data. Missing keys replaced with placeholder."""
    class SafeDict(dict):
        def __missing__(self, key):
            return f"{{{key}}}"
    return template_str.format_map(SafeDict(data))

# ------------------- Upload Route -------------------

@legal_bp.route('/upload', methods=['POST'], endpoint='upload_document')
@cross_origin()
def upload_document():
    try:
        data = request.json
        file_content = data.get('file_content')
        file_name = data.get('file_name')
        file_type = data.get('file_type')
        
        if not file_content or not file_name:
            return jsonify({'error': 'Missing file content or name'}), 400
        
        document_id = str(uuid.uuid4())
        
        try:
            decoded_content = base64.b64decode(file_content)
        except Exception:
            return jsonify({'error': 'Invalid base64 content'}), 400
        
        if file_type == 'text/plain':
            text_content = decoded_content.decode('utf-8')
        elif file_type == 'application/pdf':
            text_content = extract_pdf_text(decoded_content)
        elif file_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'application/msword']:
            text_content = extract_docx_text(decoded_content)
        else:
            text_content = "Document uploaded successfully. Content analysis available."
        
        analysis_prompt = f"""
        Analyze this legal document and provide:
        1. Document type identification
        2. Key clauses and terms
        3. Potential risks or issues
        4. Summary of main points

        Document content:
        {text_content[:2000]}
        """
        
        try:
            response = model.generate_content(analysis_prompt)
            initial_message = response.text
        except Exception:
            initial_message = f"I've received your document '{file_name}'. This appears to be a legal document. How can I assist you?"
        
        documents[document_id] = {
            'file_name': file_name,
            'file_type': file_type,
            'content': text_content,
            'analysis': initial_message
        }
        chat_histories[document_id] = []
        
        return jsonify({
            'document_id': document_id,
            'status': 'ready',
            'initial_bot_message': initial_message
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ------------------- Chat Routes -------------------

@legal_bp.route('/chat/upload', methods=['POST'], endpoint='chat_upload_document')
@cross_origin()
def chat_upload():
    try:
        data = request.json
        document_id = data.get('document_id')
        message = data.get('message')
        
        if not document_id or not message:
            return jsonify({'error': 'Missing document_id or message'}), 400
        
        if document_id not in documents:
            return jsonify({'error': 'Document not found'}), 404
        
        document = documents[document_id]
        chat_history = chat_histories.get(document_id, [])
        
        context = f"""
        Document: {document['file_name']}
        Content: {document['content'][:1500]}
        
        Previous conversation:
        """
        for chat in chat_history[-5:]:
            context += f"User: {chat['user']}\nAssistant: {chat['bot']}\n"
        
        context += f"\nUser: {message}\n\nPlease provide a helpful response based on the document content."
        
        try:
            response = model.generate_content(context)
            bot_response = response.text
        except Exception:
            bot_response = "I can help with document interpretation. Please clarify your question."
        
        chat_histories[document_id].append({'user': message, 'bot': bot_response})
        return jsonify({'bot_response': bot_response})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@legal_bp.route('/chat/general', methods=['POST'], endpoint='chat_general_document')
@cross_origin()
def chat_general():
    try:
        data = request.json
        message = data.get('message')
        
        if not message:
            return jsonify({'error': 'Missing message'}), 400
        
        prompt = f"""
        You are a legal AI assistant. Provide an informative response to this legal question:
        {message}
        """
        try:
            response = model.generate_content(prompt)
            bot_response = response.text
        except Exception:
            bot_response = "Please provide more details so I can help with your legal question."
        
        return jsonify({'bot_response': bot_response})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ------------------- Agreement Generation -------------------

@legal_bp.route('/generate-agreement', methods=['POST'], endpoint='generate_agreement_unique')
@cross_origin()
def generate_agreement_route():
    try:
        data = request.json
        agreement_type = data.get('agreement_type')
        form_data = data.get('form_data')
        
        if not agreement_type or not form_data:
            return jsonify({'error': 'Missing agreement_type or form_data'}), 400
        
        template = get_template_by_type(agreement_type)

        for section in template['sections']:
            section['title'] = safe_format(section['title'], form_data)
            section['content'] = safe_format(section['content'], form_data)
        template['signature_block'] = safe_format(template['signature_block'], form_data)

        enhanced_content = enhance_agreement_with_ai(template, agreement_type, form_data)
        pdf_path = generate_enhanced_pdf(enhanced_content, agreement_type, form_data)

        return jsonify({'pdf_url': f'/download/{os.path.basename(pdf_path)}'})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ------------------- Helper Functions for Agreement -------------------

def enhance_agreement_with_ai(template, agreement_type, form_data):
    prompt = f"""
    Review and enhance this {agreement_type} agreement template. Add missing important clauses 
    while keeping structure intact.

    Current template: {template['title']}
    Sections:
    """
    for section in template['sections']:
        prompt += f"\n{section['title']}\n{section['content']}\n"
    prompt += f"\nSignature Block:\n{template['signature_block']}"
    prompt += f"\nForm Data: {json.dumps(form_data, indent=2)}"

    try:
        response = model.generate_content(prompt)
        ai_suggestions = response.text
        enhanced_template = template.copy()
        enhanced_template['ai_suggestions'] = ai_suggestions
        return enhanced_template
    except Exception:
        return template

def generate_enhanced_pdf(template, agreement_type, form_data):
    class EnhancedPDF(FPDF):
        def __init__(self):
            super().__init__()
            self.set_auto_page_break(auto=True, margin=15)
            font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
            if os.path.exists(font_path):
                self.add_font("DejaVu", "", font_path, uni=True)
                self.add_font("DejaVu", "B", font_path, uni=True)
                self.add_font("DejaVu", "I", font_path, uni=True)
                self.set_font("DejaVu", "", 11)
            else:
                self.set_font("Arial", "", 11)
        
        def header(self):
            self.set_font('DejaVu', 'B', 20)
            self.cell(0, 15, template['title'], 0, 1, 'C')
            self.ln(5)
            self.set_draw_color(128, 128, 128)
            self.line(20, self.get_y(), 190, self.get_y())
            self.ln(10)
        
        def footer(self):
            self.set_y(-15)
            self.set_font('DejaVu', 'I', 8)
            self.set_text_color(128, 128, 128)
            self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')
        
        def add_section(self, title, content):
            self.set_font('DejaVu', 'B', 14)
            self.set_text_color(0, 0, 0)
            self.cell(0, 10, title, 0, 1, 'L')
            self.ln(2)
            self.set_font('DejaVu', '', 11)
            self.set_text_color(40, 40, 40)
            paragraphs = content.split('\n\n')
            for paragraph in paragraphs:
                if paragraph.strip():
                    lines = self.wrap_text(paragraph.strip(), 170)
                    for line in lines:
                        self.cell(0, 6, line, 0, 1, 'L')
                    self.ln(3)
            self.ln(5)
        
        def wrap_text(self, text, max_width):
            words = text.split(' ')
            lines = []
            current_line = ''
            for word in words:
                test_line = current_line + (' ' if current_line else '') + word
                if self.get_string_width(test_line) <= max_width:
                    current_line = test_line
                else:
                    if current_line:
                        lines.append(current_line)
                    current_line = word
            if current_line:
                lines.append(current_line)
            return lines
    
    pdf = EnhancedPDF()
    pdf.add_page()
    for section in template['sections']:
        filled_title = section['title'].format(**form_data)
        filled_content = section['content'].format(**form_data)
        pdf.add_section(filled_title, filled_content)
    
    pdf.ln(10)
    pdf.set_font('DejaVu', 'B', 12)
    pdf.cell(0, 10, 'SIGNATURES', 0, 1, 'L')
    pdf.ln(5)
    pdf.set_font('DejaVu', '', 11)
    signature_lines = template['signature_block'].format(**form_data).split('\n')
    for line in signature_lines:
        if line.strip():
            pdf.cell(0, 6, line, 0, 1, 'L')
    
    if 'ai_suggestions' in template:
        pdf.add_page()
        pdf.add_section('AI LEGAL REVIEW SUGGESTIONS', "Suggestions generated by AI:\n\n" + template['ai_suggestions'])
    
    temp_dir = tempfile.gettempdir()
    pdf_filename = f"{agreement_type}_{uuid.uuid4().hex[:8]}.pdf"
    pdf_path = os.path.join(temp_dir, pdf_filename)
    pdf.output(pdf_path)
    
    return pdf_path

# ------------------- Download Route -------------------

@legal_bp.route('/download/<filename>', endpoint='download_file')
@cross_origin()
def download_file(filename):
    try:
        temp_dir = tempfile.gettempdir()
        file_path = os.path.join(temp_dir, filename)
        if os.path.exists(file_path):
            return send_file(file_path, as_attachment=True, download_name=filename)
        else:
            return jsonify({'error': 'File not found'}), 404
    except Exception as e:
        return jsonify({'error': str(e)}), 500
